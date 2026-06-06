import streamlit as st
from io import BytesIO
from docx import Document
from fpdf import FPDF  # <-- Импортируем здесь, один раз


def render_extracted_text(extracted_text: str, file_info: dict):
    """Отображает извлечённый текст и кнопки скачивания"""
    st.subheader("📋 Доступные действия")

    is_large = file_info.get("size_mb", 0) > 5 or len(extracted_text) > 100000
    filename_base = file_info['name'].replace(':', '_')
    base_filename_for_export = f"{filename_base}_text"

    if is_large:
        st.info(f"📁 Текст очень большой. Для удобства предпросмотр отключён.")

        _render_download_button(
            data=extracted_text,
            filename=f"{base_filename_for_export}.md",
            key="download_original_large_md",
            label="📥 Скачать исходный текст (.md)"
        )

        _render_additional_formats(extracted_text, base_filename_for_export, key_suffix="large")

    else:
        with st.expander("📄 Показать извлечённый текст"):
            st.code(extracted_text, language="markdown")

        _render_download_button(
            data=extracted_text,
            filename=f"{base_filename_for_export}.md",
            key="download_original_small_md",
            label="📥 Скачать исходный текст (.md)"
        )

        _render_additional_formats(extracted_text, base_filename_for_export, key_suffix="small")

    st.write("---")


def _render_download_button(data: str, filename: str, key: str, label: str = "📥 Скачать"):
    """Вспомогательная функция для отрисовки кнопки скачивания"""
    st.download_button(
        label=label,
        data=data,
        file_name=filename,
        mime="text/markdown",
        key=key
    )


def _render_additional_formats(content: str, filename_base: str, key_suffix: str):
    """Отрисовывает выбор формата и кнопку скачивания для TXT, PDF, DOCX"""
    st.markdown("#### 📦 Скачать в другом формате:")

    export_format = st.selectbox(
        "Выберите формат:",
        ["TXT", "PDF", "DOCX"],
        label_visibility="collapsed",
        key=f"format_select_{key_suffix}"
    )

    if st.button("📥 Скачать выбранный формат", key=f"btn_download_new_{key_suffix}"):
        if export_format == "TXT":
            download_as_txt(content, f"{filename_base}.txt")
        elif export_format == "PDF":
            download_as_pdf(content, f"{filename_base}.pdf")
        elif export_format == "DOCX":
            download_as_docx(content, f"{filename_base}.docx")


# --- Функции экспорта ---

def download_as_txt(content: str, filename: str):
    """Скачать как TXT"""
    st.download_button(
        label="📄 Готово! Нажмите еще раз, если не скачалось",
        data=content.encode('utf-8'),
        file_name=filename,
        mime='text/plain',
        key=f"dl_txt_{filename}"
    )


def download_as_pdf(content: str, filename: str):
    """Скачать как PDF (универсальный вариант без split_text_to_size)"""
    pdf = FPDF()
    pdf.add_page()

    # Путь к системному шрифту Arial в Windows
    font_path = "C:/Windows/Fonts/arial.ttf"

    try:
        import os
        if not os.path.exists(font_path):
            raise FileNotFoundError(f"Шрифт не найден по пути: {font_path}")

        # Регистрируем шрифт с поддержкой Unicode
        pdf.add_font('Arial', '', font_path, uni=True)
        pdf.set_font('Arial', size=12)

    except Exception as e:
        st.error(f"❌ Не удалось загрузить шрифт Arial. Проверьте путь.")
        st.code(str(e))
        return

    # Простой построчный вывод без автоматического переноса длинных строк
    lines = content.split('\n')
    for line in lines:
        # Если строка очень длинная (>100 символов), разбиваем её вручную
        if len(line) > 100:
            # Разбиваем на куски по 100 символов
            chunks = [line[i:i + 100] for i in range(0, len(line), 100)]
            for chunk in chunks:
                pdf.cell(0, 10, txt=chunk, ln=True)
        else:
            pdf.cell(0, 10, txt=line, ln=True)

    buffer = BytesIO()
    pdf.output(buffer)

    st.download_button(
        label="📄 Скачать PDF",
        data=buffer.getvalue(),
        file_name=filename,
        mime='application/pdf',
        key=f"dl_pdf_{filename}"
    )


def download_as_docx(content: str, filename: str):
    """Скачать как DOCX"""
    document = Document()
    document.add_heading('Конспект Vlom', level=1)
    document.add_paragraph(content)

    buffer = BytesIO()
    document.save(buffer)

    st.download_button(
        label="📄 Готово! Нажмите еще раз, если не скачалось",
        data=buffer.getvalue(),
        file_name=filename,
        mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        key=f"dl_docx_{filename}"
    )